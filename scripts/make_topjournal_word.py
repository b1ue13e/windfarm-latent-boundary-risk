# -*- coding: utf-8 -*-
"""
Generate publication-quality Chinese Top-Journal Word (.docx) document
from manuscript/paper_chinese_topjournal_final.md
with native Microsoft Word OMML Equations, Three-Line Tables, and Embedded Figures.
"""
import re
import os
import sys
from lxml import etree
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from latex2mathml.converter import convert as latex2mathml_convert

# Find Office MML2OMML.XSL
XSLT_PATH = r'C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'
if not os.path.exists(XSLT_PATH):
    # fallback search
    for p in [
        r'C:/Program Files (x86)/Microsoft Office/root/Office16/MML2OMML.XSL',
        r'C:/Program Files/Microsoft Office/Office16/MML2OMML.XSL',
    ]:
        if os.path.exists(p):
            XSLT_PATH = p
            break

XSLT_TREE = etree.parse(XSLT_PATH) if os.path.exists(XSLT_PATH) else None
XSLT_TRANSFORM = etree.XSLT(XSLT_TREE) if XSLT_TREE else None


def clean_latex_for_mathml(eq):
    """Normalize LaTeX math for robust MathML and OMML conversion."""
    eq = eq.strip()
    eq = re.sub(r'\\tag\{[^}]+\}', '', eq).strip()
    eq = eq.replace(r'\mathrm{d}', r'd')
    eq = eq.replace(r'\Big', '')
    eq = eq.replace(r'\Bigg', '')
    eq = eq.replace(r'\big', '')
    eq = eq.replace(r'\bigg', '')
    eq = eq.replace(r'\quad', ' ')
    eq = eq.replace(r'\qquad', ' ')
    if r'\begin{cases}' in eq:
        eq = eq.replace(r'\begin{cases}', r'\left\{\begin{array}{ll}')
        eq = eq.replace(r'\end{cases}', r'\end{array}\right.')
    if r'\begin{aligned}' in eq:
        eq = eq.replace(r'\begin{aligned}', r'\begin{array}{l}')
        eq = eq.replace(r'\end{aligned}', r'\end{array}')
    return eq


def latex_to_omml_element(latex_str):
    """Convert LaTeX string to Word OMML XML element."""
    if not XSLT_TRANSFORM:
        return None
    try:
        clean_eq = clean_latex_for_mathml(latex_str)
        mathml = latex2mathml_convert(clean_eq)
        mml_tree = etree.fromstring(mathml.encode('utf-8'))
        omml_doc = XSLT_TRANSFORM(mml_tree)
        omml_root = omml_doc.getroot()
        # Parse into docx oxml element
        xml_str = etree.tostring(omml_root, encoding='utf-8')
        return parse_xml(xml_str)
    except Exception as e:
        return None


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set custom cell borders using OXML."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, val in borders.items():
        if val is not None:
            border = OxmlElement(f'w:{edge}')
            border.set(qn('w:val'), val.get('val', 'single'))
            border.set(qn('w:sz'), str(val.get('sz', 4)))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), val.get('color', 'auto'))
            tcBorders.append(border)
        else:
            border = OxmlElement(f'w:{edge}')
            border.set(qn('w:val'), 'none')
            tcBorders.append(border)
            
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Set cell inner padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_run_font(run, font_name_ascii="Times New Roman", font_name_ea="宋体", size_pt=12, bold=False, italic=False, color_rgb=(0,0,0)):
    """Set both Western and EastAsian fonts for a run."""
    run.font.name = font_name_ascii
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name_ascii)
    rFonts.set(qn('w:hAnsi'), font_name_ascii)
    rFonts.set(qn('w:eastAsia'), font_name_ea)
    rFonts.set(qn('w:cs'), font_name_ascii)
    rPr.append(rFonts)


def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=2, line_spacing=1.25, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent > 0:
        p.paragraph_format.first_line_indent = Pt(indent)
    return p


def format_mixed_text(p, text, default_size=12, default_bold=False, default_italic=False, default_ea='宋体', default_ascii='Times New Roman'):
    """Split text into bold, math ($...$), and regular text, embedding native OMML formulas for math."""
    tokens = re.split(r'(\*\*[^*]+\*\*|\$[^\$]+\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            sub_tokens = re.split(r'(\$[^\$]+\$)', inner)
            for st in sub_tokens:
                if st.startswith('$') and st.endswith('$'):
                    omml_el = latex_to_omml_element(st[1:-1])
                    if omml_el is not None:
                        p._element.append(omml_el)
                    else:
                        run = p.add_run(st[1:-1])
                        set_run_font(run, font_name_ascii="Cambria Math", font_name_ea=default_ea, size_pt=default_size, bold=True, italic=True)
                else:
                    run = p.add_run(st)
                    set_run_font(run, font_name_ascii=default_ascii, font_name_ea=default_ea, size_pt=default_size, bold=True, italic=default_italic)
        elif token.startswith('$') and token.endswith('$'):
            omml_el = latex_to_omml_element(token[1:-1])
            if omml_el is not None:
                p._element.append(omml_el)
            else:
                run = p.add_run(token[1:-1])
                set_run_font(run, font_name_ascii="Cambria Math", font_name_ea=default_ea, size_pt=default_size, bold=default_bold, italic=True)
        else:
            run = p.add_run(token)
            set_run_font(run, font_name_ascii=default_ascii, font_name_ea=default_ea, size_pt=default_size, bold=default_bold, italic=default_italic)


def render_block_equation(doc, formula_str):
    """Render a centered display block equation with right-aligned equation tag in Word."""
    raw = formula_str.strip()
    tag_m = re.search(r'\\tag\{([^}]+)\}', raw)
    tag_text = f"({tag_m.group(1)})" if tag_m else ""
    clean_eq = re.sub(r'\\tag\{[^}]+\}', '', raw).strip()

    omml_el = latex_to_omml_element(clean_eq)

    # Use a borderless 2-column table to place equation centered and tag right-aligned
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths: col 0 = 5.5 in, col 1 = 0.77 in (Total = 6.27 in matching printable area)
    table.columns[0].width = Inches(5.5)
    table.columns[1].width = Inches(0.77)

    # Cell 0: Equation centered
    cell_math = table.cell(0, 0)
    set_cell_borders(cell_math)
    set_cell_margins(cell_math, top=60, bottom=60, left=20, right=20)
    p_math = cell_math.paragraphs[0]
    p_math.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_math.paragraph_format.space_before = Pt(3)
    p_math.paragraph_format.space_after = Pt(3)
    p_math.paragraph_format.line_spacing = 1.15

    if omml_el is not None:
        p_math._element.append(omml_el)
    else:
        run = p_math.add_run(clean_eq)
        set_run_font(run, font_name_ascii="Cambria Math", font_name_ea="宋体", size_pt=11, italic=True)

    # Cell 1: Tag right-aligned
    cell_tag = table.cell(0, 1)
    set_cell_borders(cell_tag)
    set_cell_margins(cell_tag, top=60, bottom=60, left=20, right=20)
    p_tag = cell_tag.paragraphs[0]
    p_tag.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tag.paragraph_format.space_before = Pt(3)
    p_tag.paragraph_format.space_after = Pt(3)
    p_tag.paragraph_format.line_spacing = 1.15

    if tag_text:
        run_tag = p_tag.add_run(tag_text)
        set_run_font(run_tag, font_name_ascii="Times New Roman", font_name_ea="宋体", size_pt=11, bold=False)


def render_table(doc, raw_lines):
    """Render a standard Chinese academic three-line table (三线表) with native OMML cells."""
    rows_data = []
    for l in raw_lines:
        cells = [c.strip() for c in l.split('|')[1:-1]]
        if any(c.startswith(':-') or c.startswith('--') or c.endswith('-:') for c in cells):
            continue
        rows_data.append(cells)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = len(rows_data[0])

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    top_border = {'val': 'single', 'sz': 12, 'color': '000000'}    # 1.5 pt
    mid_border = {'val': 'single', 'sz': 6, 'color': '000000'}     # 0.75 pt
    bot_border = {'val': 'single', 'sz': 12, 'color': '000000'}    # 1.5 pt

    for r_idx, row in enumerate(rows_data):
        for c_idx, cell_value in enumerate(row):
            if c_idx >= num_cols:
                continue
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            is_header = (r_idx == 0)
            format_mixed_text(p, cell_value, default_size=9.5, default_bold=is_header, default_ea='黑体' if is_header else '宋体', default_ascii='Times New Roman')

            b_top = top_border if r_idx == 0 else None
            b_bot = mid_border if r_idx == 0 else (bot_border if r_idx == num_rows - 1 else None)
            set_cell_borders(cell, top=b_top, bottom=b_bot, left=None, right=None)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)


def build_docx(md_path, docx_path):
    print(f"Reading {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()

    # Set page layout to A4, 2.54 cm margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    lines = md_text.split('\n')
    idx = 0
    total_lines = len(lines)

    table_lines = []
    math_lines = []
    in_math_block = False
    in_bib_block = False

    while idx < total_lines:
        line = lines[idx].strip()

        if line == '---':
            idx += 1
            continue

        if line.startswith('```'):
            in_bib_block = not in_bib_block
            idx += 1
            continue

        if in_bib_block:
            if line.startswith('[') and ']' in line:
                p = add_styled_paragraph(doc, space_before=2, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.first_line_indent = Pt(-24)
                format_mixed_text(p, line, default_size=10, default_ea='宋体', default_ascii='Times New Roman')
            idx += 1
            continue

        # Handle Block Math ($$...$$)
        if line.startswith('$$'):
            if line.endswith('$$') and len(line) > 2:
                formula = line[2:-2].strip()
                render_block_equation(doc, formula)
                idx += 1
                continue
            else:
                in_math_block = True
                math_lines = []
                idx += 1
                continue

        if in_math_block:
            if line.endswith('$$'):
                in_math_block = False
                formula = " ".join(math_lines + [line[:-2].strip()])
                render_block_equation(doc, formula)
                idx += 1
                continue
            else:
                math_lines.append(line)
                idx += 1
                continue

        # Handle Tables (| ... |)
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            idx += 1
            if idx < total_lines and not (lines[idx].strip().startswith('|') and lines[idx].strip().endswith('|')):
                render_table(doc, table_lines)
                table_lines = []
            continue

        # Handle Images (![...](...))
        m_img = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if m_img:
            img_rel_path = m_img.group(2).strip()
            cand_paths = [
                img_rel_path,
                os.path.abspath(img_rel_path),
                os.path.join(os.path.dirname(os.path.abspath(md_path)), img_rel_path),
                os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(md_path))), img_rel_path),
                os.path.join("figures", os.path.basename(img_rel_path)),
            ]
            found_img = None
            for cp in cand_paths:
                if os.path.exists(cp):
                    found_img = cp
                    break

            if found_img:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.keep_with_next = True
                run_img = p_img.add_run()
                run_img.add_picture(found_img, width=Inches(5.8))
                print(f"  [Word] Embedded figure: {found_img}")
            else:
                print(f"  [Word WARNING] Figure not found: {img_rel_path}")

            idx += 1
            continue

        # Handle Figure Caption (**图 X** ...)
        if line.startswith('**图') and ('**' in line[3:]):
            p = add_styled_paragraph(doc, space_before=2, space_after=8, line_spacing=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_mixed_text(p, line, default_size=10.5, default_bold=True, default_ea='黑体', default_ascii='Times New Roman')
            idx += 1
            continue

        # Handle Table Caption (**表 X** ...)
        if line.startswith('**表') and ('**' in line[3:]):
            p = add_styled_paragraph(doc, space_before=10, space_after=4, line_spacing=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
            format_mixed_text(p, line, default_size=10.5, default_bold=True, default_ea='黑体', default_ascii='Times New Roman')
            idx += 1
            continue

        # Handle Title (# ...)
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = add_styled_paragraph(doc, space_before=14, space_after=12, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)
            run = p.add_run(title_text)
            set_run_font(run, font_name_ascii="SimHei", font_name_ea="黑体", size_pt=20, bold=True)
            idx += 1
            continue

        # Handle Level 1 Heading (## ...)
        if line.startswith('## '):
            h_text = line[3:].strip()
            p = add_styled_paragraph(doc, space_before=14, space_after=6, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            set_run_font(run, font_name_ascii="SimHei", font_name_ea="黑体", size_pt=15, bold=True)
            idx += 1
            continue

        # Handle Level 2 Heading (### ...)
        if line.startswith('### '):
            h_text = line[4:].strip()
            if h_text == '摘要':
                p = add_styled_paragraph(doc, space_before=10, space_after=4, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)
                p.paragraph_format.keep_with_next = True
                run = p.add_run("摘  要")
                set_run_font(run, font_name_ascii="SimHei", font_name_ea="黑体", size_pt=13, bold=True)
            else:
                p = add_styled_paragraph(doc, space_before=10, space_after=4, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(h_text)
                set_run_font(run, font_name_ascii="SimHei", font_name_ea="黑体", size_pt=13, bold=True)
            idx += 1
            continue

        # Handle Author block
        if line.startswith('**作者信息**') or line.startswith('**作者单位**') or line.startswith('**通讯作者**'):
            p = add_styled_paragraph(doc, space_before=2, space_after=4, line_spacing=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_mixed_text(p, line, default_size=11, default_ea='楷体', default_ascii='Times New Roman')
            idx += 1
            continue

        # Handle Keywords
        if line.startswith('**关键词**'):
            p = add_styled_paragraph(doc, space_before=6, space_after=14, line_spacing=1.25, indent=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.right_indent = Pt(24)
            format_mixed_text(p, line, default_size=10.5, default_ea='楷体', default_ascii='Times New Roman')
            idx += 1
            continue

        if not line:
            idx += 1
            continue

        # Regular body paragraph
        p = add_styled_paragraph(doc, space_before=0, space_after=3, line_spacing=1.25, indent=24, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        format_mixed_text(p, line, default_size=12, default_ea='宋体', default_ascii='Times New Roman')
        idx += 1

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"Successfully saved Word document to: {docx_path}")


if __name__ == '__main__':
    src = "manuscript/paper_chinese_topjournal_final.md"
    dst = "manuscript/论文_潜伏控制相变与遥测退化约束下的风电运行备用决策_最小充分模型准则.docx"
    build_docx(src, dst)
    dst2 = "build/paper_chinese_topjournal.docx"
    build_docx(src, dst2)
